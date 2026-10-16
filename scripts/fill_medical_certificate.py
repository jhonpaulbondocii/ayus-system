#!/usr/bin/env python3
"""
scripts/fill_medical_certificate.py

Fills the MEDICAL-CERTIFICATE-ISO-2025.docx template with provided data.

Usage:
  python3 scripts/fill_medical_certificate.py \
    --template public/template/MEDICAL-CERTIFICATE-ISO-2025.docx \
    --data /tmp/cert-data-xxx.json \
    --output /tmp/cert-out-xxx.docx

JSON data keys:
  date           - e.g. "October 14, 2025"
  honorific      - "Mr." | "Ms." | "Mrs."
  name           - full name of student
  age_sex        - e.g. "20/M"
  course         - e.g. "Bachelor of Science in Information Technology"
  date_examined  - e.g. "October 14, 2025"
  diagnosis      - e.g. "Upper Respiratory Tract Infection (URTI)"
  recommendation - e.g. "rest for 3 days and to take prescribed medications"
  ptr_number     - e.g. "2025-12345" (empty string if none)
"""

import argparse
import json
import sys
from docx import Document


def set_para_text(para, new_text: str):
    """
    Replace all runs in a paragraph with a single run containing new_text.
    Preserves the formatting of the first run.
    """
    if not para.runs:
        return
    para.runs[0].text = new_text
    for run in para.runs[1:]:
        run.text = ""


def fill_certificate(template_path: str, data: dict, output_path: str):
    doc = Document(template_path)
    paras = doc.paragraphs

    # P1: "Date: \t" → "Date: <date>"
    if len(paras) > 1:
        set_para_text(paras[1], f"Date: {data.get('date', '')}")

    # P7: the main body paragraph with tab placeholders
    # Original: "This is to certify that Mr./Ms./Mrs. \t\t\t  age/sex \t, of \t\t,
    #            was seen and examined on \tand was diagnosed to have \t\t\t\t
    #            I therefore recommend \t\t\t "
    if len(paras) > 7:
        honorific      = data.get("honorific", "Mr./Ms./Mrs.")
        name           = data.get("name", "")
        age_sex        = data.get("age_sex", "")
        course         = data.get("course", "")
        date_examined  = data.get("date_examined", "")
        diagnosis      = data.get("diagnosis", "")
        recommendation = data.get("recommendation", "")

        body = (
            f"This is to certify that {honorific} {name}  age/sex {age_sex}, "
            f"of {course}, was seen and examined on {date_examined} "
            f"and was diagnosed to have {diagnosis} "
            f"I therefore recommend {recommendation}"
        )
        set_para_text(paras[7], body)

    # P15: "University Physician License No. 70611 PTR #: \t"
    if len(paras) > 15:
        ptr = data.get("ptr_number", "")
        ptr_text = f"University Physician License No. 70611 PTR #: {ptr}" if ptr else \
                   "University Physician License No. 70611"
        set_para_text(paras[15], ptr_text)

    doc.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Fill medical certificate template")
    parser.add_argument("--template", required=True, help="Path to the .docx template")
    parser.add_argument("--data",     required=True, help="Path to JSON data file")
    parser.add_argument("--output",   required=True, help="Output .docx path")
    args = parser.parse_args()

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    fill_certificate(args.template, data, args.output)
    print(f"Certificate written to {args.output}", file=sys.stderr)


if __name__ == "__main__":
    main()